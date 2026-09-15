"""Orquestador del pipeline de ingesta de documentos.

ingest_document -> detect_type -> [ocr] -> extract_text -> chunk -> embed -> guardar

Contrato: `specs/sprint-05-rag.md` §2. Lo invoca
`app/tasks/document_ingestion.py::_run_pipeline` (Dev B, ya en `main`), que
tambien es quien marca el documento como `failed` si algo de aqui lanza — por
eso `process()` no atrapa excepciones: las deja subir tal cual, para que el
worker decida si reintenta o lo manda a la Dead Letter Queue.

Cada paso queda en su propia transaccion (`tenant_session()` no admite commits
intermedios: es una sola transaccion por bloque `async with`), asi que si el
worker muere entre "processing" y "completed", el documento no se queda en un
estado sin persistir: como minimo, "processing" ya esta commiteado.
"""

import csv
import io
import logging
from typing import Any
from uuid import UUID

import pymupdf
from docx import Document as DocxFile
from openpyxl import load_workbook

from app.core.database import tenant_session
from app.models.document import Document
from app.models.document_chunk import DocumentChunk
from app.services.chunker import DocumentChunker
from app.services.embedding import EmbeddingService
from app.services.ocr import OCRService
from app.services.storage import download_from_storage

logger = logging.getLogger(__name__)

# Tipos que siempre necesitan OCR: no tienen capa de texto que extraer.
_IMAGE_FILE_TYPES = ("png", "jpg", "webp")

# Si un PDF "de texto" trae menos caracteres que esto en total, probablemente es
# un escaneo sin capa de texto (o con una capa vacia/rota) y hay que pasarlo por OCR.
_PDF_MIN_TEXT_CHARS = 50


class EmptyDocumentError(ValueError):
    """El documento no produjo ningun chunk de texto (extraccion u OCR vacios).

    Un escaneo ilegible o un archivo corrupto no deben quedar como `completed`
    con `chunk_count=0`: eso es indistinguible de un documento sin contenido
    relevante, y RAG nunca podria citarlo. `app/tasks/document_ingestion.py` la
    trata como `PipelineUnavailableError`: no se reintenta (el mismo archivo
    va a dar el mismo resultado) y el documento queda `failed` con el motivo
    visible en `metadata.error`.
    """


def extract_pdf_pages(content: bytes) -> list[dict[str, Any]]:
    """Extrae el texto nativo de cada pagina de un PDF (sin OCR).

    Args:
        content: Bytes del PDF.

    Returns:
        Una entrada `{"page_number": int, "text": str}` por pagina.
    """
    with pymupdf.open(stream=content, filetype="pdf") as documento:
        return [
            {"page_number": i, "text": pagina.get_text()}
            for i, pagina in enumerate(documento, start=1)
        ]


def extract_docx_text(content: bytes) -> list[dict[str, Any]]:
    """Extrae el texto de un DOCX como una sola "pagina" (Word no pagina).

    Args:
        content: Bytes del archivo `.docx`.

    Returns:
        Una unica entrada con `page_number=None` (no aplica) y todo el texto.
    """
    documento = DocxFile(io.BytesIO(content))
    texto = "\n".join(parrafo.text for parrafo in documento.paragraphs)
    return [{"page_number": None, "text": texto}]


def extract_csv_text(content: bytes) -> list[dict[str, Any]]:
    """Convierte un CSV a texto tabular simple, separado por tabs.

    Args:
        content: Bytes del archivo `.csv`.

    Returns:
        Una unica entrada con todas las filas.
    """
    texto_csv = content.decode("utf-8", errors="replace")
    filas = csv.reader(io.StringIO(texto_csv))
    texto = "\n".join("\t".join(fila) for fila in filas)
    return [{"page_number": None, "text": texto}]


def extract_xlsx_text(content: bytes) -> list[dict[str, Any]]:
    """Convierte cada hoja de un XLSX a texto tabular, una "pagina" por hoja.

    Args:
        content: Bytes del archivo `.xlsx`.

    Returns:
        Una entrada por hoja, con su indice (1-based) como `page_number`.
    """
    libro = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        paginas = []
        for i, hoja in enumerate(libro.worksheets, start=1):
            filas = (
                "\t".join("" if celda is None else str(celda) for celda in fila)
                for fila in hoja.iter_rows(values_only=True)
            )
            paginas.append({"page_number": i, "text": "\n".join(filas)})
        return paginas
    finally:
        libro.close()


# Extractores de texto sin OCR, por extension normalizada (ver ALLOWED_TYPES en
# app/api/v1/documents.py). Las imagenes no tienen entrada: siempre van por OCR.
_TEXT_EXTRACTORS = {
    "pdf": extract_pdf_pages,
    "docx": extract_docx_text,
    "csv": extract_csv_text,
    "xlsx": extract_xlsx_text,
}


class DocumentPipeline:
    """Ejecuta el pipeline completo de ingesta para un documento.

    Attributes:
        chunker: Servicio de chunking.
        embedder: Servicio de generacion de embeddings.
        ocr: Servicio de OCR para imagenes y PDFs escaneados.
    """

    def __init__(
        self,
        chunker: DocumentChunker,
        embedder: EmbeddingService,
        ocr: OCRService,
    ) -> None:
        """Ensambla el pipeline con sus dependencias.

        Args:
            chunker: Servicio de chunking a usar.
            embedder: Servicio de embeddings a usar.
            ocr: Servicio de OCR a usar.
        """
        self.chunker = chunker
        self.embedder = embedder
        self.ocr = ocr

    async def process(self, document_id: UUID, client_id: UUID) -> None:
        """Procesa un documento de punta a punta: descarga, extrae, chunkea, embebe.

        Args:
            document_id: Documento a procesar.
            client_id: Tenant propietario.

        Raises:
            ValueError: Si el documento no existe (para este tenant o para
                ninguno: desaparecio entre el encolado y el procesamiento).
            EmptyDocumentError: Si la extraccion (u OCR) no produjo texto
                utilizable en ninguna pagina y por lo tanto no hay nada que
                chunkear ni embeber.
        """
        file_url, file_type, title = await self._marcar_procesando(document_id, client_id)

        contenido = await download_from_storage(file_url)

        if self._needs_ocr(file_type, contenido):
            text_pages = await self.ocr.extract_text(contenido, file_type)
        else:
            text_pages = self._extract_text(contenido, file_type)

        chunks = self.chunker.chunk(
            text_pages=text_pages,
            document_title=title,
            file_type=file_type,
        )

        if not chunks:
            raise EmptyDocumentError(
                f"El documento {document_id} no genero ningun chunk de texto "
                "(extraccion u OCR vacios)"
            )

        embeddings = await self.embedder.embed_batch(
            texts=[chunk.content for chunk in chunks],
            batch_size=100,
        )

        await self._guardar_resultado(document_id, client_id, title, chunks, embeddings)

        logger.info(
            "Documento %s procesado: %s chunks",
            document_id,
            len(chunks),
            extra={"client_id": str(client_id)},
        )

    async def _marcar_procesando(self, document_id: UUID, client_id: UUID) -> tuple[str, str, str]:
        """Carga el documento, lo marca `processing` y devuelve lo que hace falta despues.

        Se leen los campos aqui (no se devuelve el objeto ORM) porque la sesion
        se cierra al salir de `tenant_session()`: acceder a atributos de una
        instancia con la sesion ya cerrada es fragil, leerlos como valores
        simples no lo es.

        Args:
            document_id: Documento a procesar.
            client_id: Tenant propietario.

        Returns:
            `(file_url, file_type, title)` del documento.

        Raises:
            ValueError: Si el documento no existe para este tenant.
        """
        async with tenant_session(client_id) as session:
            document = await session.get(Document, document_id)
            if document is None:
                raise ValueError(f"Documento {document_id} no encontrado")

            document.status = "processing"
            await session.flush()

            return document.file_url or "", document.file_type or "", document.title

    async def _guardar_resultado(
        self,
        document_id: UUID,
        client_id: UUID,
        title: str,
        chunks: list[Any],
        embeddings: list[list[float]],
    ) -> None:
        """Guarda los chunks generados y marca el documento como `completed`.

        Args:
            document_id: Documento procesado.
            client_id: Tenant propietario.
            title: Titulo del documento, para la metadata de cada chunk.
            chunks: Chunks generados por el chunker.
            embeddings: Un embedding por chunk, en el mismo orden.

        Raises:
            ValueError: Si el documento desaparecio durante el procesamiento.
        """
        async with tenant_session(client_id) as session:
            document = await session.get(Document, document_id)
            if document is None:
                raise ValueError(f"Documento {document_id} desaparecio durante el procesamiento")

            for i, (chunk, embedding) in enumerate(zip(chunks, embeddings, strict=True)):
                session.add(
                    DocumentChunk(
                        client_id=client_id,
                        document_id=document_id,
                        chunk_index=i,
                        content=chunk.content,
                        embedding=embedding,
                        token_count=chunk.token_count,
                        metadata_={
                            "page_number": chunk.page_number,
                            "section": chunk.section,
                            "source_file": title,
                        },
                    )
                )

            document.status = "completed"
            document.chunk_count = len(chunks)

    def _needs_ocr(self, file_type: str, content: bytes) -> bool:
        """Decide si un archivo necesita pasar por OCR.

        Las imagenes siempre lo necesitan. Un PDF lo necesita si su capa de
        texto esta vacia o es sospechosamente corta (probable escaneo sin OCR
        previo, o con una capa de texto rota).

        Args:
            file_type: Extension normalizada del archivo.
            content: Bytes del archivo.

        Returns:
            True si hay que pasarlo por `OCRService`.
        """
        if file_type in _IMAGE_FILE_TYPES:
            return True

        if file_type == "pdf":
            try:
                paginas = extract_pdf_pages(content)
            except Exception:
                logger.warning("No se pudo leer la capa de texto del PDF; se intenta OCR")
                return True
            total_texto = "".join(pagina["text"] for pagina in paginas)
            return len(total_texto.strip()) < _PDF_MIN_TEXT_CHARS

        return False

    def _extract_text(self, content: bytes, file_type: str) -> list[dict[str, Any]]:
        """Extrae texto de un archivo que no necesita OCR.

        Args:
            content: Bytes del archivo.
            file_type: Extension normalizada del archivo.

        Returns:
            Paginas extraidas, en el formato que consume el chunker.

        Raises:
            ValueError: Si `file_type` no tiene extractor (p. ej. una imagen:
                esas siempre pasan por OCR, nunca llegan aqui).
        """
        extractor = _TEXT_EXTRACTORS.get(file_type)
        if extractor is None:
            raise ValueError(f"Tipo de archivo no soportado para extraccion: {file_type}")
        return extractor(content)
